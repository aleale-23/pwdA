import openpyxl
from openpyxl import Workbook


def lookup_terms(
    file_master_path: str,
    file_alamat_path: str,
    col_relasi_master: int = 3,
    col_terms_master: int = 7,
    col_relasi_alamat: int = 5,
) -> list:
    """FUNGSI 1: LOOKUP MASTER KE TEMPLATE ALAMAT"""
    idx_relasi_m = col_relasi_master - 1
    idx_terms_m = col_terms_master - 1
    idx_relasi_a = col_relasi_alamat - 1

    wb_master = openpyxl.load_workbook(file_master_path, data_only=True)
    ws_master = wb_master.active

    master_dict = {}
    for row in ws_master.iter_rows(min_row=2, values_only=True):
        if len(row) > max(idx_relasi_m, idx_terms_m):
            kode = row[idx_relasi_m]
            terms = row[idx_terms_m]
            if kode is not None:
                master_dict[str(kode).strip()] = terms

    wb_alamat = openpyxl.load_workbook(file_alamat_path, data_only=True)
    ws_alamat = wb_alamat.active

    header_alamat = [cell.value for cell in ws_alamat[1]]
    data_lengkap = [list(header_alamat) + ["terms"]]

    for row in ws_alamat.iter_rows(min_row=2, values_only=True):
        if len(row) > idx_relasi_a:
            kode = row[idx_relasi_a]
            kode_bersih = str(kode).strip() if kode is not None else ""
            terms_val = master_dict.get(kode_bersih, None)
            data_lengkap.append(list(row) + [terms_val])

    return data_lengkap


def filter_bukan_cash_dan_na(data_berisi_terms: list) -> list:
    """FUNGSI 2: FILTER NILAI TERMS"""
    header = data_berisi_terms[0]
    idx_terms = len(header) - 1

    kata_abaikan = {"cash and carry", "n/a", "na", "#n/a", "none", ""}
    data_terfilter = [header]

    for row in data_berisi_terms[1:]:
        terms_val = row[idx_terms]
        if terms_val is None:
            continue
        if str(terms_val).strip().lower() in kata_abaikan:
            continue
        data_terfilter.append(row)

    return data_terfilter


def hapus_duplikat_kode_relasi(data: list, col_relasi: int = 5) -> list:
    """FUNGSI 3: HAPUS DUPLIKAT BERDASARKAN KODE RELASI"""
    idx_relasi = col_relasi - 1
    header = data[0]
    data_unik = [header]

    kode_tercatat = set()

    for row in data[1:]:
        kode = row[idx_relasi]
        kode_bersih = str(kode).strip() if kode is not None else ""

        if not kode_bersih or kode_bersih in kode_tercatat:
            continue

        kode_tercatat.add(kode_bersih)
        data_unik.append(row)

    return data_unik


def gabung_dengan_invoice(
    data_sebelumnya: list,
    file_invoice_path: str,
    col_relasi_invoice: int = 1,
    col_relasi_target: int = 5,
) -> list:
    """FUNGSI 4: MERGE / APPEND DARI INVOICE_TAGIH"""
    wb_inv = openpyxl.load_workbook(file_invoice_path, data_only=True)
    ws_inv = wb_inv.active

    idx_inv = col_relasi_invoice - 1
    idx_target = col_relasi_target - 1

    header = data_sebelumnya[0]
    total_kolom = len(header)

    data_gabungan = [row for row in data_sebelumnya]

    for row in ws_inv.iter_rows(min_row=2, values_only=True):
        if len(row) > idx_inv:
            kode = row[idx_inv]
            if kode is None:
                continue

            baris_baru = [None] * total_kolom
            baris_baru[idx_target] = str(kode).strip()
            baris_baru[-1] = "DARI INVOICE"

            data_gabungan.append(baris_baru)

    return data_gabungan


def lookup_detail_alamat(
    data_unik: list,
    file_master_alamat_path: str,
    col_relasi_sumber: int = 5,
    col_relasi_target: int = 2,
    col_alamat: int = 4,
    col_kota: int = 5,
    col_telepon: int = 6,
) -> list:
    """FUNGSI 5: LOOKUP KE MASTER DATA ALAMAT

    Menghasilkan 4 kolom: [kode_relasi, alamat, kota, nomor_telepon]
    """
    idx_sumber = col_relasi_sumber - 1
    idx_relasi_t = col_relasi_target - 1
    idx_alamat = col_alamat - 1
    idx_kota = col_kota - 1
    idx_telp = col_telepon - 1

    # 1. Baca master_data_alamat ke dictionary {kode_relasi: (alamat, kota, telp)}
    wb_m_alamat = openpyxl.load_workbook(
        file_master_alamat_path, data_only=True
    )
    ws_m_alamat = wb_m_alamat.active

    dict_alamat = {}
    for row in ws_m_alamat.iter_rows(min_row=2, values_only=True):
        if len(row) > max(idx_relasi_t, idx_alamat, idx_kota, idx_telp):
            kode = row[idx_relasi_t]
            if kode is not None:
                clean_kode = str(kode).strip()
                alamat = row[idx_alamat]
                kota = row[idx_kota]
                telp = row[idx_telp]
                dict_alamat[clean_kode] = (alamat, kota, telp)

    # 2. Susun tabel hasil akhir 4 kolom
    header_final = ["kode_relasi", "alamat", "kota", "nomor_telepon"]
    hasil_final = [header_final]

    for row in data_unik[1:]:
        kode = row[idx_sumber]
        clean_kode = str(kode).strip() if kode is not None else ""

        # Cari data alamat dari dictionary
        detail = dict_alamat.get(clean_kode, (None, None, None))

        baris_final = [clean_kode, detail[0], detail[1], detail[2]]
        hasil_final.append(baris_final)

    return hasil_final


def simpan_ke_excel(data: list, output_path: str):
    """FUNGSI 6: SIMPAN KE FILE EXCEL"""
    wb = Workbook()
    ws = wb.active
    ws.title = "Hasil Akhir"

    for baris in data:
        ws.append(baris)

    wb.save(output_path)
    print(f"File berhasil disimpan ke: {output_path}")


# --- ALUR UTAMA ---
if __name__ == "__main__":
    file_master = "master_data.xlsx"
    file_alamat = "template_alamat.xlsx"
    file_invoice = "invoice_tagih.xlsx"
    file_master_alamat = "master_data_alamat.xlsx"
    file_output = "hasil_final_alamat_lengkap.xlsx"

    # Posisi kolom (1-based)
    KOLOM_RELASI_MASTER = 3
    KOLOM_TERMS_MASTER = 7
    KOLOM_RELASI_ALAMAT = 5
    KOLOM_RELASI_INVOICE = 1

    # 1. Lookup file alamat ke master data terms
    data_lookup = lookup_terms(
        file_master_path=file_master,
        file_alamat_path=file_alamat,
        col_relasi_master=KOLOM_RELASI_MASTER,
        col_terms_master=KOLOM_TERMS_MASTER,
        col_relasi_alamat=KOLOM_RELASI_ALAMAT,
    )
    print(f"1. Selesai Lookup Terms            : {len(data_lookup) - 1} baris")

    # 2. Filter nilai terms (buang cash and carry & N/A)
    data_filter = filter_bukan_cash_dan_na(data_lookup)
    print(f"2. Selesai Filter Terms            : {len(data_filter) - 1} baris")

    # 3. Hapus duplikat data alamat
    data_unik_alamat = hapus_duplikat_kode_relasi(
        data_filter, col_relasi=KOLOM_RELASI_ALAMAT
    )
    print(
        f"3. Selesai Hapus Duplikat Alamat   : {len(data_unik_alamat) - 1} baris"
    )

    # 4. Gabungkan dengan file invoice_tagih
    data_gabung_inv = gabung_dengan_invoice(
        data_sebelumnya=data_unik_alamat,
        file_invoice_path=file_invoice,
        col_relasi_invoice=KOLOM_RELASI_INVOICE,
        col_relasi_target=KOLOM_RELASI_ALAMAT,
    )
    print(
        f"4. Selesai Merge Invoice           : {len(data_gabung_inv) - 1} baris"
    )

    # 5. Hapus duplikat hasil gabungan
    data_final_unik = hapus_duplikat_kode_relasi(
        data_gabung_inv, col_relasi=KOLOM_RELASI_ALAMAT
    )
    print(
        f"5. Selesai Hapus Duplikat Akhir    : {len(data_final_unik) - 1} baris"
    )

    # 6. Lookup ke file master_data_alamat (ambil alamat, kota, telepon)
    output_4_kolom = lookup_detail_alamat(
        data_unik=data_final_unik,
        file_master_alamat_path=file_master_alamat,
        col_relasi_sumber=KOLOM_RELASI_ALAMAT,  # kolom 5
        col_relasi_target=2,  # kolom 2 di master_data_alamat
        col_alamat=4,  # kolom 4 di master_data_alamat
        col_kota=5,  # kolom 5 di master_data_alamat
        col_telepon=6,  # kolom 6 di master_data_alamat
    )
    print(
        f"6. Selesai Lookup Detail Alamat    : {len(output_4_kolom) - 1} baris"
    )

    # 7. Simpan hasil akhir (4 kolom)
    simpan_ke_excel(output_4_kolom, file_output)
